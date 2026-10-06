#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
export-brief-docx.py — 把周报 Markdown（总览）按 Jack 的 Word 模版导出为 .docx

用法（每周复用）：
    python3 export-brief-docx.py \
        --md   /path/to/weekly-brief-YYYY-MM-DD.md \
        --date "2026 年 9 月 24 日" \
        --out  "/path/to/Vehicle - 美欧日汽车政策情报 20260924.docx" \
        [--template /path/to/vehicle-template-v1.1.docx] \
        [--title "美欧日汽车政策情报"] \
        [--version V1.0] \
        [--start "## 一句话判断"] [--end "## 待证实项"] \
        [--pdf]            # 顺带用 soffice 转 PDF（同目录、同名 .pdf）

行为：
  * 以模版 docx 为底（保留 styles.xml、页眉页脚、页面设置、主题、编号定义），清空正文后写入新内容。
  * 只导出 Markdown 中从 --start（默认 "## 一句话判断"）开始、到 --end（默认 "## 待证实项"）之前的内容；
    文件顶部的标题/链接列表和 --end 之后的内容不导出。
  * 标题区：模版 Title 样式（居中，Arial Unicode MS）+ 灰色居中日期行；--version 非空时以 "日期 · V1.0" 形式写入日期行。
  * "## " → 模版 heading 1（自动加"一、二、…"）；"### " → 模版 heading 2（自动加"（一）（二）…"）。
  * 正文/列表项：模版正文写法（Normal + Arial Unicode MS 16pt）；列表用项目符号编号（新增一个 bullet 定义）。
  * `[[信源]](url)` → 正文后的小号灰色超链接，显示为 "[信源]"；普通 `[文字](url)` → 模版 Hyperlink 样式超链接。
  * `**粗体**` 保留；Markdown 表格 → Word 表格（模版无表格样式，用与模版风格接近的细线简洁表格）。
  * 末尾追加模版的 "（完）"。

依赖：python-docx、lxml（pip install python-docx lxml）；--pdf 需要 LibreOffice（soffice）。
"""
from __future__ import annotations

import argparse
import copy
import os
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE.parent / "templates" / "vehicle-template-v1.1.docx"

# ---- 模版实测样式要点（见 styles.xml / document.xml）-----------------------------
TPL_FONT = "Arial Unicode MS"      # 模版所有正文/标题 run 均显式指定该字体
BODY_SZ = Pt(16)                   # 正文 run sz=32（16pt），Normal 样式本身为宋体 12pt
H1_SZ = Pt(16)                     # heading 1 样式 24pt，但模版 run 覆盖为 16pt 粗体
SUBTITLE_COLOR = RGBColor(0x5A, 0x5A, 0x5A)   # 模版日期/依据行颜色
SOURCE_COLOR = RGBColor(0x80, 0x80, 0x80)     # 信源链接：小号灰色
SOURCE_SZ = Pt(9)
TABLE_SZ = Pt(10.5)
TABLE_HEADER_FILL = "E7E6E6"
TABLE_BORDER_COLOR = "808080"

STYLE_TITLE = "Title"
STYLE_H1 = "Heading 1"
STYLE_H2 = "Heading 2"
STYLE_HYPERLINK = "Hyperlink"

CN_NUM = "一二三四五六七八九十"


def cn_ordinal(n: int) -> str:
    if n <= 10:
        return CN_NUM[n - 1]
    if n < 20:
        return "十" + CN_NUM[n - 11]
    return str(n)


# ---- Markdown 解析 -------------------------------------------------------------
def slice_markdown(text: str, start: str, end: str | None) -> list[str]:
    lines = text.splitlines()
    try:
        i0 = next(i for i, l in enumerate(lines) if l.strip().startswith(start))
    except StopIteration:
        sys.exit(f"找不到起始标记：{start!r}")
    i1 = len(lines)
    if end:
        for i in range(i0 + 1, len(lines)):
            if lines[i].strip().startswith(end):
                i1 = i
                break
    return lines[i0:i1]


# inline tokens: source tag [[label]](url), link [text](url), bold **x**
INLINE_RE = re.compile(
    r"(?P<source>\[\[(?P<slabel>[^\]]+)\]\]\((?P<surl>[^)\s]+)\))"
    r"|(?P<link>\[(?P<llabel>[^\]]+)\]\((?P<lurl>[^)\s]+)\))"
    r"|(?P<bold>\*\*(?P<btext>.+?)\*\*)"
    r"|(?P<code>`(?P<ctext>[^`]+)`)"
)


def smart_quotes(text: str) -> str:
    """把直引号 "..." 成对转成中文弯引号 “...”（模版用法），单引号同理。"""
    out = []
    open_d = True
    open_s = True
    for ch in text:
        if ch == '"':
            out.append("“" if open_d else "”")
            open_d = not open_d
        elif ch == "'":
            out.append("‘" if open_s else "’")
            open_s = not open_s
        else:
            out.append(ch)
    return "".join(out)


def parse_inline(text: str) -> list[tuple[str, str, str | None]]:
    """返回 [(kind, text, url)]，kind ∈ text|bold|source|link"""
    out: list[tuple[str, str, str | None]] = []
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            out.append(("text", smart_quotes(text[pos:m.start()]), None))
        if m.group("source"):
            # 标签内用不换行空格，避免 "[Reuters/Straits Times]" 被拆到两行
            label = m.group("slabel").replace(" ", "\u00a0")
            out.append(("source", f"[{label}]", m.group("surl")))
        elif m.group("link"):
            out.append(("link", smart_quotes(m.group("llabel")), m.group("lurl")))
        elif m.group("bold"):
            out.append(("bold", smart_quotes(m.group("btext")), None))
        else:
            out.append(("text", m.group("ctext"), None))
        pos = m.end()
    if pos < len(text):
        out.append(("text", smart_quotes(text[pos:]), None))
    # 信源标签前面紧贴正文时补一个空格，并合并连续空白
    cleaned: list[tuple[str, str, str | None]] = []
    for kind, t, url in out:
        if kind == "text":
            t = re.sub(r"[ \t]{2,}", " ", t)
            if not t:
                continue
        cleaned.append((kind, t, url))
    return cleaned


def parse_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for l in lines:
        l = l.strip()
        if not l.startswith("|"):
            continue
        cells = [c.strip() for c in l.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue  # 分隔行
        rows.append(cells)
    return rows


# ---- docx 写入辅助 ---------------------------------------------------------------
def set_run_font(run, size=None, bold=None, color=None, font=TPL_FONT):
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font)
    if size is not None:
        run.font.size = size
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    # 模版 run 写的是 w:lang val="zh-CN"，这会让 LibreOffice 对西文也按中文规则任意断行；
    # 这里改为 eastAsia="zh-CN"，Word 中外观不变，且西文单词/数字不再被拆开。
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "zh-CN")


def add_hyperlink(paragraph, text: str, url: str, *, small_gray: bool, size=BODY_SZ):
    part = paragraph.part
    r_id = part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True
    )
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rstyle = OxmlElement("w:rStyle")
    rstyle.set(qn("w:val"), paragraph.part.document.styles[STYLE_HYPERLINK].style_id)
    rpr.append(rstyle)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    hl.append(new_run)
    paragraph._p.append(hl)
    from docx.text.run import Run

    run = Run(new_run, paragraph)
    if small_gray:
        set_run_font(run, size=SOURCE_SZ, color=SOURCE_COLOR)
        run.font.underline = False
    else:
        set_run_font(run, size=size)
    return run


def write_inline(paragraph, text: str, size=BODY_SZ, bold_base=False):
    tokens = parse_inline(text)
    for idx, (kind, t, url) in enumerate(tokens):
        if kind in ("source", "link"):
            prev_is_text = idx > 0 and tokens[idx - 1][0] in ("text", "bold")
            if kind == "source" and prev_is_text and not tokens[idx - 1][1].endswith((" ", "\u3000")):
                r = paragraph.add_run(" ")
                set_run_font(r, size=size)
            add_hyperlink(paragraph, t, url, small_gray=(kind == "source"), size=size)
        else:
            if kind == "text" and idx > 0 and tokens[idx - 1][0] == "source" and not t.startswith(" "):
                t = " " + t
            r = paragraph.add_run(t)
            set_run_font(r, size=size, bold=(True if kind == "bold" else (bold_base or None)))


def clear_body(doc):
    body = doc.element.body
    for el in list(body):
        if el.tag != qn("w:sectPr"):
            body.remove(el)


def ensure_bullet_num(doc) -> int:
    """在模版 numbering.xml 中新增一个 '•' 项目符号定义，返回 numId。"""
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(a.get(qn("w:abstractNumId"))) for a in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
    new_abs = max(abstract_ids, default=-1) + 1
    new_num = max(num_ids, default=0) + 1

    abs_el = OxmlElement("w:abstractNum")
    abs_el.set(qn("w:abstractNumId"), str(new_abs))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "singleLevel")
    abs_el.append(mlt)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, val in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
        e = OxmlElement(tag)
        e.set(qn("w:val"), val)
        lvl.append(e)
    ppr = OxmlElement("w:pPr")
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "420")
    ind.set(qn("w:hanging"), "420")
    ppr.append(ind)
    lvl.append(ppr)
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), TPL_FONT)
    rpr.append(rf)
    lvl.append(rpr)
    abs_el.append(lvl)

    # abstractNum 必须排在所有 num 之前
    nums = numbering.findall(qn("w:num"))
    if nums:
        nums[0].addprevious(abs_el)
    else:
        numbering.append(abs_el)
    num_el = OxmlElement("w:num")
    num_el.set(qn("w:numId"), str(new_num))
    an = OxmlElement("w:abstractNumId")
    an.set(qn("w:val"), str(new_abs))
    num_el.append(an)
    numbering.append(num_el)
    return new_num


def set_list_item(paragraph, num_id: int):
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl)
    numpr.append(nid)
    ppr.append(numpr)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    ppr.append(spacing)


def add_title_block(doc, title: str, date_line: str, version: str | None):
    p = doc.add_paragraph(style=STYLE_TITLE)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(title)
    set_run_font(r)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(12)
    txt = date_line if not version else f"{date_line} · {version}"
    r2 = p2.add_run(txt)
    set_run_font(r2, color=SUBTITLE_COLOR)


def add_heading1(doc, text: str, n: int):
    p = doc.add_paragraph(style=STYLE_H1)
    r = p.add_run(f"{cn_ordinal(n)}、{smart_quotes(text)}")
    set_run_font(r, size=H1_SZ)
    return p


def add_heading2(doc, text: str, n: int):
    p = doc.add_paragraph(style=STYLE_H2)
    r = p.add_run(f"（{cn_ordinal(n)}）{smart_quotes(text)}")
    set_run_font(r)
    return p


def add_body(doc, text: str):
    p = doc.add_paragraph()
    write_inline(p, text)
    return p


def add_bullet(doc, text: str, num_id: int):
    p = doc.add_paragraph()
    set_list_item(p, num_id)
    write_inline(p, text)
    return p


def _set_cell_borders(tbl):
    tblpr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), TABLE_BORDER_COLOR)
        borders.append(e)
    tblpr.append(borders)


def _shade(cell, fill: str):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=ncols)
    tbl.autofit = False
    _set_cell_borders(tbl)
    # 列宽（twips）：模版文本区宽 11906-1800*2 = 8306；三列表按 日期/区域/事件 分配，其它列数均分
    if ncols == 3:
        widths = (2000, 850, 5456)
    else:
        widths = tuple([8306 // ncols] * ncols)
    tblpr = tbl._tbl.tblPr
    tblw = tblpr.find(qn("w:tblW"))
    if tblw is None:
        tblw = OxmlElement("w:tblW")
        tblpr.append(tblw)
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    for gc, w in zip(tbl._tbl.tblGrid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = tbl.cell(ri, ci)
            text = row[ci] if ci < len(row) else ""
            p = cell.paragraphs[0]
            ppr = p._p.get_or_add_pPr()
            snap = OxmlElement("w:snapToGrid")   # 表内不吸附行网格，行高紧凑
            snap.set(qn("w:val"), "0")
            ppr.append(snap)
            p.paragraph_format.space_before = Pt(2.5)
            p.paragraph_format.space_after = Pt(2.5)
            write_inline(p, text, size=TABLE_SZ, bold_base=(ri == 0))
            if ri == 0:
                _shade(cell, TABLE_HEADER_FILL)
            tcw = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tcw is None:
                tcw = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tcw)
            tcw.set(qn("w:w"), str(widths[ci]))
            tcw.set(qn("w:type"), "dxa")
    # 表头行跨页重复
    trpr = tbl.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    trpr.append(th)
    return tbl


# ---- 主流程 ------------------------------------------------------------------------
def build(md_path: Path, date_line: str, out_path: Path, template: Path, title: str, version: str | None,
          start: str, end: str | None):
    text = md_path.read_text(encoding="utf-8")
    lines = slice_markdown(text, start, end)

    doc = Document(str(template))
    clear_body(doc)
    bullet_num = ensure_bullet_num(doc)
    add_title_block(doc, title, date_line, version)

    h1 = h2 = 0
    i = 0
    para_buf: list[str] = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            add_body(doc, " ".join(s.strip() for s in para_buf))
            para_buf = []

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush_para()
            i += 1
            continue
        if s.startswith("## "):
            flush_para()
            h1 += 1
            h2 = 0
            add_heading1(doc, s[3:].strip(), h1)
        elif s.startswith("### "):
            flush_para()
            h2 += 1
            add_heading2(doc, s[4:].strip(), h2)
        elif s.startswith("#"):
            flush_para()
            add_body(doc, s.lstrip("#").strip())
        elif s.startswith("|"):
            flush_para()
            j = i
            while j < len(lines) and lines[j].strip().startswith("|"):
                j += 1
            add_table(doc, parse_table(lines[i:j]))
            i = j
            continue
        elif re.match(r"^[-*+]\s+", s) or re.match(r"^\d+[.)]\s+", s):
            flush_para()
            item = re.sub(r"^([-*+]|\d+[.)])\s+", "", s)
            # 续行（缩进的非列表行）并入本项
            j = i + 1
            while j < len(lines) and lines[j].startswith((" ", "\t")) and lines[j].strip() \
                    and not re.match(r"^\s*([-*+]|\d+[.)])\s+", lines[j]):
                item += " " + lines[j].strip()
                j += 1
            add_bullet(doc, item, bullet_num)
            i = j
            continue
        else:
            para_buf.append(s)
        i += 1
    flush_para()

    end_p = doc.add_paragraph()
    end_p.paragraph_format.space_before = Pt(12)
    r = end_p.add_run("（完）")
    set_run_font(r, size=BODY_SZ)

    # 文档属性
    doc.core_properties.title = f"{title} {date_line}"
    doc.core_properties.subject = ""
    doc.core_properties.keywords = ""

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path


def to_pdf(docx_path: Path) -> Path | None:
    try:
        subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
            check=True, capture_output=True, timeout=180,
        )
    except Exception as e:  # noqa: BLE001
        print(f"[warn] PDF 转换失败：{e}", file=sys.stderr)
        return None
    return docx_path.with_suffix(".pdf")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--md", required=True, type=Path, help="输入 Markdown（周报总览）")
    ap.add_argument("--date", required=True, help='日期行文字，如 "2026 年 9 月 24 日"')
    ap.add_argument("--out", required=True, type=Path, help="输出 .docx 路径")
    ap.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    ap.add_argument("--title", default="美欧日汽车政策情报")
    ap.add_argument("--version", default="", help='版本号，如 V1.0；为空则日期行不带版本')
    ap.add_argument("--start", default="## 一句话判断", help="导出起点（行前缀匹配）")
    ap.add_argument("--end", default="## 待证实项", help="导出终点（不含该节）；传空字符串则到文末")
    ap.add_argument("--pdf", action="store_true", help="同时用 LibreOffice 转 PDF")
    args = ap.parse_args()

    out = build(args.md, args.date, args.out, args.template, args.title, args.version or None,
                args.start, args.end or None)
    print(f"docx: {out}")
    if args.pdf:
        pdf = to_pdf(out)
        if pdf:
            print(f"pdf:  {pdf}")


if __name__ == "__main__":
    main()
